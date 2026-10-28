import importlib.util
import json
from pathlib import Path
import stat
import tempfile
import unittest
import zipfile

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location("advanced",ROOT/"src-tauri/python/advanced_tools.py")
advanced=importlib.util.module_from_spec(spec);spec.loader.exec_module(advanced)

class AdvancedToolsTest(unittest.TestCase):
    def setUp(self): self.tmp=tempfile.TemporaryDirectory();self.path=Path(self.tmp.name)
    def tearDown(self): self.tmp.cleanup()
    def call(self,operation,source,output,**kwargs):
        return advanced.main(dict(operation=operation,input_paths=[str(source)],output_path=str(output),native_helper=str(ROOT/"src-tauri/native-bin/vert-document-tools"),**kwargs))
    def test_structured_roundtrip_preserves_types_and_unicode(self):
        data={"中文":[1,True,None,{"quote":"a,b\n\"quoted\""}],"text":"<x>&"};source=self.path/"input.json";source.write_text(json.dumps(data))
        for ext in ["yaml","xml","json"]:
            target=self.path/("roundtrip."+ext);self.call("structured",source,target);source=target
        self.assertEqual(json.loads(source.read_text()),data)
    def test_xml_mixed_content_roundtrip(self):
        source=self.path/"source.xml";source.write_text('<root a="1">before<x>inside</x>after<x>second</x></root>')
        data=self.path/"data.json";result=self.path/"result.xml";self.call("structured",source,data);self.call("structured",data,result)
        self.assertEqual(advanced.read_data(source),advanced.read_data(result))
    def test_csv_roundtrip_preserves_quotes_multiline_and_duplicate_headers(self):
        source=self.path/"input.csv";source.write_text('Name,Name\n中文,"a,b"\n"Line\nbreak","He said ""Hi"""\n')
        value=advanced.read_data(source);data=self.path/"data.json";target=self.path/"out.csv";self.call("structured",source,data);self.call("structured",data,target)
        self.assertEqual(advanced.read_data(target),value)
    def test_json_object_not_silently_flattened_to_csv(self):
        source=self.path/"input.json";source.write_text('{"nested":{"a":1}}')
        with self.assertRaisesRegex(ValueError,"表格结构"):self.call("structured",source,self.path/"out.csv")
    def test_html_table_to_csv_preserves_values(self):
        source=self.path/"table.html";source.write_text('<html><table><tr><th>姓名</th><th>备注</th></tr><tr><td>中文</td><td>a,b</td></tr></table></html>')
        target=self.path/"out.csv";self.call("structured",source,target)
        self.assertEqual(advanced.read_data(target),{"columns":["姓名","备注"],"rows":[["中文","a,b"]]})
    def test_xml_entities_and_yaml_cycles_rejected(self):
        source=self.path/"input.xml";source.write_text('<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>')
        with self.assertRaises(Exception):advanced.read_data(source)
        source=self.path/"cycle.yaml";source.write_text('a: &a [*a]')
        with self.assertRaisesRegex(ValueError,"循环"):advanced.read_data(source)
    def test_archive_rejects_path_traversal_symlinks_and_duplicates(self):
        for index,kind in enumerate(["traversal","symlink","duplicate"]):
            archive=self.path/f"{index}.zip"
            with zipfile.ZipFile(archive,"w") as out:
                if kind=="traversal":out.writestr("../escape.txt","bad")
                elif kind=="symlink":
                    info=zipfile.ZipInfo("link");info.create_system=3;info.external_attr=(stat.S_IFLNK|0o777)<<16;out.writestr(info,"/tmp/target")
                else:out.writestr("A.txt","one");out.writestr("a.txt","two")
            with self.assertRaises(ValueError):self.call("extract-zip",archive,self.path/f"extract{index}")
        self.assertFalse((self.path/"escape.txt").exists())
    def test_zip_creation_duplicates_and_nested_extraction(self):
        first=self.path/"a";second=self.path/"b";first.mkdir();second.mkdir();(first/"中文.txt").write_text("first");(second/"中文.txt").write_text("second")
        archive=self.path/"out.zip";advanced.main(dict(operation="create-zip",input_paths=[str(first/"中文.txt"),str(second/"中文.txt")],output_path=str(archive)))
        self.call("extract-zip",archive,self.path/"extracted")
        self.assertEqual((self.path/"extracted/中文.txt").read_text(),"first");self.assertEqual((self.path/"extracted/中文 (1).txt").read_text(),"second")
    def test_pdf_word_html_and_scanned_fallback(self):
        from docx import Document
        source=ROOT/"work/fixtures/digital.pdf";output=self.path/"document.docx"
        self.call("pdf-office",source,output);doc=Document(output);text="\n".join(p.text for p in doc.paragraphs)+"\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn("VERT Enhanced",text);self.assertIn("中文",text)
        html=self.path/"document.html";self.call("pdf-office",source,html);self.assertIn("VERT Enhanced",html.read_text())
        output=self.path/"scan.docx";_,warnings=self.call("pdf-office",ROOT/"work/fixtures/scanned.pdf",output)
        self.assertTrue(any("OCR" in warning for warning in warnings));self.assertIn("中文","\n".join(p.text for p in Document(output).paragraphs))
    def test_pdf_blank_page_is_preserved_without_ocr_error(self):
        import pymupdf as fitz
        from docx import Document
        doc=fitz.open(ROOT/"work/fixtures/digital.pdf");doc.new_page();source=self.path/"blank-page.pdf";doc.save(source)
        target=self.path/"blank-page.docx";self.call("pdf-office",source,target)
        output=Document(target);self.assertIn("VERT Enhanced","\n".join(p.text for p in output.paragraphs))
        self.assertGreaterEqual(output.element.xml.count('w:type="page"'),3)
    def test_native_jpeg_page_range_is_applied(self):
        import subprocess
        output=self.path/"jpeg-pages"
        request=dict(operation="pdf-images",input_paths=[str(ROOT/"work/fixtures/digital.pdf")],output_path=str(output),pages="2",image_format="jpg")
        result=subprocess.run([str(ROOT/"src-tauri/native-bin/vert-document-tools"),json.dumps(request)],capture_output=True,text=True)
        self.assertEqual(result.returncode,0,result.stderr);self.assertEqual([p.name for p in output.iterdir()],["page-0002.jpg"])
    def table_fixture(self):
        from reportlab.pdfgen import canvas
        path=self.path/"table.pdf";c=canvas.Canvas(str(path),pagesize=(612,792))
        rows=[["Name","Value","Unit"],["Alpha","123","kg"],["Beta","456","kg"],["=1+1","789","kg"]]
        xs=[40,220,370,550];ys=[700,670,640,610,580]
        for x in xs:c.line(x,580,x,700)
        for y in ys:c.line(40,y,550,y)
        for row_index,row in enumerate(rows):
            for column,text in enumerate(row):c.drawString(xs[column]+8,ys[row_index]-20,text)
        c.save();return path
    def test_pdf_tables_export_actual_cells_and_disable_formulas(self):
        from openpyxl import load_workbook
        target=self.path/"table.xlsx";self.call("pdf-office",self.table_fixture(),target);book=load_workbook(target)
        table=book["P1-T1"];values=[c.value for row in table for c in row];self.assertIn(123,values);self.assertIn("Alpha",values)
        formula=next(c for row in table for c in row if c.value=="=1+1");self.assertEqual(formula.data_type,"s")
        with self.assertRaisesRegex(ValueError,"未找到"):self.call("pdf-office",ROOT/"work/fixtures/scanned.pdf",self.path/"scan.xlsx")
    def test_pdf_password_decryption(self):
        from pypdf import PdfReader,PdfWriter
        source=ROOT/"work/fixtures/digital.pdf";writer=PdfWriter(clone_from=source);writer.encrypt("test-secret",algorithm="AES-256");encrypted=self.path/"encrypted.pdf";writer.write(encrypted)
        with self.assertRaisesRegex(ValueError,"密码不正确"):self.call("decrypt-pdf",encrypted,self.path/"wrong.pdf",password="incorrect")
        target=self.path/"unlocked.pdf";self.call("decrypt-pdf",encrypted,target,password="test-secret");reader=PdfReader(target);self.assertFalse(reader.is_encrypted);self.assertEqual(len(reader.pages),3)

if __name__=="__main__":unittest.main(verbosity=2)
