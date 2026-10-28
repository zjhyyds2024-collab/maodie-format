"""Local adapters. Requests live in a private file; passwords never enter argv."""
from __future__ import annotations
import base64
import csv
import html
import io
import json
import os
import re
from pathlib import Path, PurePosixPath
import shutil
import stat
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET

# Local modules must not write bytecode into the signed application bundle.
sys.dont_write_bytecode = True

MAX_MEMBERS = 10000
MAX_EXTRACT_BYTES = 2 * 1024**3
MAX_TEXT_BYTES = 64 * 1024**2

def emit(value):
    print(json.dumps(value, ensure_ascii=False), flush=True)

def progress(current, total, stage):
    emit(dict(current=current, total=total, stage=stage))

def read_text(path):
    data = Path(path).read_bytes()
    if len(data) > MAX_TEXT_BYTES:
        raise ValueError("文本文件超过 64 MB 处理上限。")
    for encoding in (["utf-16"] if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "gb18030"]):
        try:
            return data.decode(encoding)
        except UnicodeError:
            pass
    raise ValueError("无法识别文本编码，请先保存为 UTF-8。")

def validate_data(value, ancestors=None, budget=None):
    ancestors = set() if ancestors is None else ancestors
    budget = [0] if budget is None else budget
    budget[0] += 1
    if budget[0] > 500000:
        raise ValueError("结构化文本的节点数量过多。")
    if isinstance(value, (dict, list)):
        if id(value) in ancestors:
            raise ValueError("不支持循环引用的 YAML。")
        ancestors.add(id(value))
        if isinstance(value, dict):
            for key, item in value.items():
                if not isinstance(key, str):
                    raise ValueError("JSON/XML 互转要求对象键为字符串。")
                validate_data(item, ancestors, budget)
        else:
            for item in value:
                validate_data(item, ancestors, budget)
        ancestors.remove(id(value))
    elif not isinstance(value, (str, int, float, bool, type(None))):
        raise ValueError("包含不适合通用格式转换的 YAML 类型，请使用字符串、数字、数组和对象。")
    json.dumps(value, ensure_ascii=False, allow_nan=False)

def xml_tree(node):
    return dict(tag=node.tag, attributes=node.attrib, text=node.text, tail=node.tail, children=[xml_tree(child) for child in node])

def rebuild_tree(value):
    node = ET.Element(value["tag"], value.get("attributes", {}))
    node.text, node.tail = value.get("text"), value.get("tail")
    for child in value.get("children", []):
        node.append(rebuild_tree(child))
    return node

def typed_xml(value, name="value"):
    node = ET.Element(name)
    if isinstance(value, dict):
        node.set("type", "object")
        for key, item in value.items():
            child = typed_xml(item, "item"); child.set("key", key); node.append(child)
    elif isinstance(value, list):
        node.set("type", "array")
        for item in value: node.append(typed_xml(item, "item"))
    else:
        node.set("type", "scalar"); node.text = json.dumps(value, ensure_ascii=False)
    return node

def from_typed_xml(node):
    if node.get("type") == "object": return {child.attrib["key"]: from_typed_xml(child) for child in node}
    if node.get("type") == "array": return [from_typed_xml(child) for child in node]
    return json.loads(node.text or "null")

def read_data(path):
    import yaml
    from defusedxml import ElementTree as SafeET
    ext = Path(path).suffix.lower()
    text = read_text(path)
    if ext == ".json": value = json.loads(text)
    elif ext in (".yaml", ".yml"): value = yaml.safe_load(text)
    elif ext == ".xml":
        from defusedxml.common import DefusedXmlException
        try: root = SafeET.fromstring(text)
        except DefusedXmlException as error:
            raise ValueError("XML 包含不安全的实体或外部引用，已拒绝处理。") from error
        value = from_typed_xml(root) if root.get("format") == "vert-structured-v1" else {"xml_tree": xml_tree(root)}
    elif ext == ".csv":
        rows = list(csv.reader(io.StringIO(text, newline="")))
        value = {"columns": rows[0] if rows else [], "rows": rows[1:]}
    else: value = {"text": text}
    validate_data(value)
    return value

def structured(request):
    import yaml
    source, output = Path(request["input_paths"][0]), Path(request["output_path"])
    value = read_data(source); target = output.suffix.lower(); warnings = []
    if source.suffix.lower() in (".html", ".htm") and target == ".csv":
        from lxml import html as lhtml
        tree=lhtml.fromstring(read_text(source),parser=lhtml.HTMLParser(no_network=True))
        tables=tree.xpath("//table")
        if not tables: raise ValueError("HTML 中未找到可导出的表格，请选择含表格的网页。")
        if tables:
            rows=[[cell.text_content().strip() for cell in row.xpath("./th|./td")] for row in tables[0].xpath("./tr|./thead/tr|./tbody/tr|./tfoot/tr")]
            rows=[row for row in rows if row]
            if not rows: raise ValueError("HTML 表格为空，没有可导出的单元格。")
            value={"columns":rows[0],"rows":rows[1:]}
            if len(tables)>1:warnings.append("HTML 包含多张表格，CSV 仅导出第一张表格。")
    if target == ".json": text = json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)
    elif target in (".yaml", ".yml"): text = yaml.safe_dump(value, allow_unicode=True, sort_keys=False)
    elif target == ".xml":
        root = rebuild_tree(value["xml_tree"]) if isinstance(value, dict) and set(value) == {"xml_tree"} else typed_xml(value, "document")
        if not (isinstance(value, dict) and set(value) == {"xml_tree"}): root.set("format", "vert-structured-v1")
        text = ET.tostring(root, encoding="unicode", xml_declaration=True)
    elif target == ".csv":
        if isinstance(value, dict) and isinstance(value.get("columns"), list) and isinstance(value.get("rows"), list): rows = [value["columns"], *value["rows"]]
        elif isinstance(value, list) and all(isinstance(row, dict) for row in value):
            keys = list(dict.fromkeys(key for row in value for key in row)); rows = [keys] + [[row.get(key, "") for key in keys] for row in value]
        elif isinstance(value, list) and all(isinstance(row, list) for row in value): rows = value
        elif isinstance(value, dict) and set(value) == {"text"}: rows = [["text"], [value["text"]]]
        else: raise ValueError("转 CSV 需要表格结构：对象数组、二维数组，或 columns/rows 对象；不能将任意嵌套数据无损转成表格。")
        stream = io.StringIO(newline=""); writer = csv.writer(stream)
        for row in rows:
            writer.writerow([json.dumps(item, ensure_ascii=False) if isinstance(item, (dict, list)) else item for item in row])
        text = "\ufeff" + stream.getvalue()
    elif target == ".html":
        if isinstance(value, dict) and set(value) == {"columns", "rows"}:
            body = "<table border='1'>" + "".join("<tr>" + "".join("<td>" + html.escape(str(cell)) + "</td>" for cell in row) + "</tr>" for row in [value["columns"], *value["rows"]]) + "</table>"
        else: body = "<pre><code>" + html.escape(value["text"] if isinstance(value, dict) and set(value)=={"text"} else json.dumps(value, ensure_ascii=False, indent=2)) + "</code></pre>"
        text = '<!doctype html><html lang="zh"><meta charset="utf-8"><body>' + body + '</body></html>'
    elif target == ".md" and isinstance(value, dict) and set(value)=={"columns","rows"}:
        def cell(item): return str(item).replace("|", "\\|").replace("\n", "<br>")
        rows=[value["columns"], ["---"]*len(value["columns"]), *value["rows"]]
        text="\n".join("| " + " | ".join(cell(item) for item in row) + " |" for row in rows)
    else: text = value["text"] if isinstance(value, dict) and set(value) == {"text"} else json.dumps(value, ensure_ascii=False, indent=2)
    output.write_text(text, encoding="utf-8")
    return [str(output)], warnings

def safe_extract(source, destination):
    destination = Path(destination); destination.mkdir(parents=True, exist_ok=True)
    names = set(); total = 0; files = []
    with zipfile.ZipFile(source) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_MEMBERS: raise ValueError("ZIP 文件条目超过 10000 个。")
        # Validate the complete central directory before writing any member.
        for member in entries:
            name = member.filename.replace("\\", "/"); path = PurePosixPath(name)
            mode = member.external_attr >> 16
            if path.is_absolute() or ".." in path.parts or not path.parts or ":" in path.parts[0] or "\x00" in name:
                raise ValueError("ZIP 包含越界路径，已拒绝解压。")
            if stat.S_ISLNK(mode) or (stat.S_IFMT(mode) and not (stat.S_ISREG(mode) or stat.S_ISDIR(mode))):
                raise ValueError("ZIP 包含符号链接或特殊文件，已拒绝解压。")
            identity = str(path).casefold()
            if identity in names: raise ValueError("ZIP 包含重复或大小写冲突的路径。")
            names.add(identity); total += member.file_size
            if member.flag_bits & 1: raise ValueError("暂不支持带密码的 ZIP。")
            if total > MAX_EXTRACT_BYTES or member.file_size > MAX_EXTRACT_BYTES: raise ValueError("ZIP 展开大小超过 2 GB。")
        written = 0
        for index, member in enumerate(entries):
            name = member.filename.replace("\\", "/"); target = destination.joinpath(*PurePosixPath(name).parts)
            if member.is_dir(): target.mkdir(parents=True, exist_ok=True); continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with archive.open(member) as source_file, target.open("xb") as output:
                while block := source_file.read(1024 * 1024):
                    written += len(block)
                    if written > MAX_EXTRACT_BYTES: raise ValueError("ZIP 实际展开大小超过上限。")
                    output.write(block)
            files.append(str(target)); progress(index+1, len(entries), "正在解压")
    return files

def zip_files(paths, output):
    seen = set()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        for index, path in enumerate(paths):
            path = Path(path)
            if not path.is_file() or path.is_symlink(): raise ValueError("压缩入口只接受普通文件。")
            name = path.name; stem, suffix = path.stem, path.suffix; n = 1
            while name.casefold() in seen: name = f"{stem} ({n}){suffix}"; n += 1
            seen.add(name.casefold()); archive.write(path, name); progress(index+1,len(paths),"正在压缩")
    return [str(output)], []

def normalize_cjk(text):
    import unicodedata
    return "".join(unicodedata.normalize("NFKC", char) if 0x2F00 <= ord(char) <= 0x2FDF or 0xF900 <= ord(char) <= 0xFAFF else char for char in text)

def pdf_table_value(value, force_text=False):
    """Infer only simple, short numbers; retain identifiers, dates and formula-like text."""
    if value is None: return ""
    original = str(value); text = original.strip()
    if not force_text and re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?", text):
        digits = sum(char.isdigit() for char in text)
        if digits <= 15:
            return float(text) if "." in text else int(text)
    return original

def identifier_header(value):
    text = str(value or "").strip().lower()
    return bool(re.search(r"编号|编码|账号|账户|电话|手机|邮编|证件|身份证|学号|代码|订单|\b(?:id|code|account|phone|zip|postal)\b", text))

def pdf_to_office(request):
    import pymupdf as fitz
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.styles import Font
    source, output = Path(request["input_paths"][0]), Path(request["output_path"])
    warnings = []; doc = fitz.open(source)
    if doc.needs_pass: raise ValueError("PDF 已加密，请先使用密码解密工具。")
    if len(doc) == 0: raise ValueError("PDF 不包含页面。")
    if output.suffix == ".xlsx":
        import pdfplumber
        book = Workbook(); book.remove(book.active); report = book.create_sheet("处理说明")
        report.append(["页码", "检测到的表格", "说明"]); count = 0
        with pdfplumber.open(source) as pdf:
            for index, page in enumerate(pdf.pages):
                tables = page.extract_tables()
                if not tables:
                    candidate = page.extract_tables({"vertical_strategy":"text", "horizontal_strategy":"text", "min_words_vertical":3})
                    # Text alignment alone is not proof of a table. Require several rows and columns.
                    tables = [table for table in candidate if len(table)>=3 and sum(sum(bool(str(c or "").strip()) for c in row)>=2 for row in table)>=3]
                for table_index, table in enumerate(tables):
                    if not any(any(str(cell or "").strip() for cell in row) for row in table): continue
                    count += 1; sheet = book.create_sheet(f"P{index+1}-T{table_index+1}")
                    protected = [identifier_header(value) for value in table[0]]
                    for row_index, row in enumerate(table):
                        values = [pdf_table_value(value, force_text=row_index == 0 or (index < len(protected) and protected[index])) for index, value in enumerate(row)]
                        sheet.append(values)
                        for cell in sheet[sheet.max_row]:
                            # Explicit strings prevent formula activation without disabling real numbers.
                            if isinstance(cell.value, str): cell.data_type = "s"
                    sheet.freeze_panes="A2"
                    for column in sheet.columns: sheet.column_dimensions[column[0].column_letter].width=22
                detail = "原生表格提取；明确的短数字按数值写入，编号/日期/公式文本保留；请核对列对应" if tables else "未检测到可靠表格；此页未导出"
                report.append([index+1,len(tables),detail]); progress(index+1,len(pdf.pages),"正在提取 PDF 表格")
                if not tables: warnings.append(f"第 {index+1} 页未检测到可靠表格，未作为表格导出。")
        if not count: raise ValueError("未找到可提取的原生表格；扫描表格和复杂表格暂不能可靠转 Excel，请勿将普通文字当作表格。")
        report["A1"].font=Font(bold=True); book.save(output)
        warnings.append("PDF 表格提取为实验功能：明确且不超过 15 位的数字按数值写入，编号、日期和公式保留文本；请检查列和处理说明。")
        return [str(output)], warnings
    texts = [normalize_cjk(page.get_text("text", sort=True)) for page in doc]
    blank_pages = [not text.strip() and not page.get_images() and not page.get_drawings() for text, page in zip(texts, doc)]
    has_scan = any(not text.strip() and not blank for text, blank in zip(texts, blank_pages))
    if output.suffix == ".docx" and not has_scan and not any(blank_pages):
        from pdf2docx import Converter
        converter = Converter(str(source))
        conversion_ok = True
        try: converter.convert(str(output), multi_processing=False)
        except Exception:
            conversion_ok = False
            warnings.append("版式转换失败，已改为逐页文字重建。")
        finally: converter.close()
        exported = Document(output) if conversion_ok else Document()
        paragraphs = list(exported.paragraphs) + [p for table in exported.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
        for paragraph in paragraphs:
            for run in paragraph.runs: run.text = normalize_cjk(run.text)
        if conversion_ok: exported.save(output)
        extracted = "\n".join([p.text for p in exported.paragraphs]+[cell.text for table in exported.tables for row in table.rows for cell in row.cells])
        original = "".join(texts)
        from collections import Counter
        original_chars=Counter(c for c in original if not c.isspace()); result_chars=Counter(c for c in extracted if not c.isspace())
        coverage=sum((original_chars & result_chars).values())/max(1,sum(original_chars.values()))
        if coverage >= .95:
            warnings.append("已按 PDF 版式转换为可编辑 Word；复杂分栏、图片位置和字体仍需人工核对。")
            return [str(output)], warnings
        warnings.append("版式转换文字覆盖不足，已改为逐页文字重建，保留可提取文字。")
    helper = request.get("native_helper")
    output_doc = Document() if output.suffix == ".docx" else None
    html_pages = []
    for index, text in enumerate(texts):
        if not text.strip() and not blank_pages[index]:
            if not helper: raise ValueError("扫描页需要本地 OCR 组件。")
            import subprocess
            with tempfile.TemporaryDirectory(dir=output.parent) as tmp:
                textfile=Path(tmp)/"page.txt"
                response=subprocess.run([helper,json.dumps(dict(operation="ocr",input_paths=[str(source)],output_path=str(textfile),pages=str(index+1)))],capture_output=True,text=True)
                if response.returncode: raise ValueError(f"第 {index+1} 页 OCR 失败：{response.stderr.strip()[:300]}")
                text=textfile.read_text(); warnings.append(f"第 {index+1} 页使用 OCR，未恢复原始版式。")
        if output_doc is not None:
            if index: output_doc.add_page_break()
            for line in text.splitlines(): output_doc.add_paragraph(line)
        else: html_pages.append(f'<section><h2>第 {index+1} 页</h2><pre>{html.escape(text)}</pre></section>')
        progress(index+1,len(texts),"正在重建文档")
    if output_doc is not None: output_doc.save(output)
    else: output.write_text('<!doctype html><html lang="zh"><meta charset="utf-8"><title>PDF text</title><body>'+"\n".join(html_pages)+'</body></html>',encoding="utf-8")
    warnings.append("此结果为可编辑文字重建，不保证原 PDF 排版；图片内容请使用 PDF 转图片。")
    return [str(output)], warnings


def pdf_page_selection(text, count, required=False):
    import re
    text = str(text or '').strip().replace('，', ',').replace('–', '-')
    if not text:
        if required: raise ValueError('请填写要处理的页码。')
        return list(range(count))
    if len(text) > 20000: raise ValueError('页码输入过长。')
    result = []; seen = set()
    for part in text.split(','):
        part = part.strip()
        if not re.fullmatch(r'[0-9]+(?:-[0-9]+)?', part): raise ValueError('页码格式应为 1-3,5；页码从 1 开始。')
        values = part.split('-'); start = int(values[0]); end = int(values[-1])
        if start < 1 or end < start or end > count: raise ValueError(f'页码超出文档范围（共 {count} 页）。')
        for index in range(start - 1, end):
            if index in seen: raise ValueError('页码不能重复，请检查页序。')
            seen.add(index); result.append(index)
    return result

def pdf_page_tools(request):
    from pypdf import PdfReader, PdfWriter
    source = Path(request['input_paths'][0]); output = Path(request['output_path'])
    if source.resolve() == output.resolve(): raise ValueError('请另存为副本，不能覆盖源 PDF。')
    reader = PdfReader(source)
    if reader.is_encrypted: raise ValueError('PDF 已加密，请先使用密码解密工具。')
    count = len(reader.pages)
    if count < 1 or count > 10000: raise ValueError('页面整理支持 1–10000 页的 PDF。')
    operation = request['operation']
    selected = pdf_page_selection(request.get('pages'), count, operation != 'rotate-pdf')
    warnings = []
    if operation == 'rotate-pdf':
        angle = request.get('rotation', 90)
        if type(angle) is not int or angle not in {90, 180, 270}: raise ValueError('旋转角度应为 90、180 或 270。')
        writer = PdfWriter(clone_from=reader)
        for position, index in enumerate(selected):
            writer.pages[index].rotate(angle); progress(position + 1, len(selected), '正在旋转页面')
    else:
        if operation == 'delete-pages':
            excluded = set(selected)
            selected = [index for index in range(count) if index not in excluded]
            if not selected: raise ValueError('不能删除全部页面，请至少保留一页。')
        elif operation == 'reorder-pages':
            if len(selected) != count: raise ValueError('重排需包含每一页且仅出现一次，例如 3,1,2。')
        elif operation != 'extract-pages': raise ValueError('不支持的页面操作。')
        writer = PdfWriter()
        for position, index in enumerate(selected):
            writer.add_page(reader.pages[index]); progress(position + 1, len(selected), '正在整理页面')
        if reader.metadata: writer.add_metadata({str(key): value for key, value in reader.metadata.items() if isinstance(value, str)})
        root = reader.trailer['/Root']
        if any(key in root for key in ['/Outlines', '/AcroForm', '/Names']) or any(page.get('/Annots') for page in reader.pages):
            warnings.append('页面已整理；原文件包含书签、表单、附件或注释，请核对这些内容和页面链接。')
    writer.write(output)
    return [str(output)], warnings

def decrypt_pdf(request):
    from pypdf import PdfReader, PdfWriter
    source, output=Path(request["input_paths"][0]),Path(request["output_path"])
    reader=PdfReader(source)
    if reader.is_encrypted and not reader.decrypt(request.get("password") or ""):
        raise ValueError("密码不正确，PDF 未解密。")
    writer=PdfWriter(clone_from=reader); writer.write(output)
    return [str(output)], [] if reader.is_encrypted else ["该 PDF 未加密，已另存为副本。"]

def main(request):
    output=Path(request["output_path"]);output.parent.mkdir(parents=True,exist_ok=True)
    operation=request["operation"]
    if operation=="structured": return structured(request)
    if operation in {"rotate-pdf", "extract-pages", "delete-pages", "reorder-pages"}: return pdf_page_tools(request)
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    if operation == "music-preview":
        from encrypted_music import music_preview
        return music_preview(request)
    if operation in {"decode-music", "music-cover"}:
        from encrypted_music import decode_music, music_cover
        return decode_music(request, progress) if operation == "decode-music" else music_cover(request)
    if operation=="decode-ncm":
        from ncm_tools import decode_ncm
        return decode_ncm(request, progress)
    from archive_tools import CREATE_FORMATS, create_archive, extract_archive
    if operation in CREATE_FORMATS: return create_archive(operation, request["input_paths"], output, progress)
    if operation=="create-zip": return zip_files(request["input_paths"],output)
    if operation=="extract-zip":
        source = request["input_paths"][0]
        return (safe_extract(source,output) if Path(source).suffix.lower()==".zip" else extract_archive(source,output,progress)), []
    if operation=="pdf-office": return pdf_to_office(request)
    if operation=="decrypt-pdf": return decrypt_pdf(request)
    raise ValueError("不支持的扩展文档操作。")

if __name__=="__main__":
    try:
        request=json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
        outputs,warnings=main(request);emit(dict(outputs=outputs,warnings=warnings))
    except Exception as error:
        print(str(error),file=sys.stderr);sys.exit(1)
